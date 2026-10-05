#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成批量二维码工具的测试文件和示例输出。"""
import csv
from pathlib import Path
from openpyxl import Workbook

from qr_generator import render_qr

def make_normal_roster(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["CodeID", "Content"])
    ws.append(["Q001", "https://example.com/course/python"])
    ws.append(["Q002", "WIFI:T:WPA;S:TrainingRoom;P:12345678;;"])
    ws.append(["Q003", "mailto:training@example.com"])
    ws.append(["Q004", "TEL:+8613800000000"])
    ws.append(["Q005", "欢迎参加2026年秋季培训"])
    wb.save(path)

def make_abnormal_roster(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["CodeID", "Content"])
    ws.append(["Q001", "https://example.com/a"])
    ws.append(["q001", "https://example.com/b"])   # 与 Q001 重复（忽略大小写）
    ws.append(["Q003", ""])                          # Content 为空
    ws.append(["", "https://example.com/d"])        # CodeID 为空
    ws.append(["A/01", "https://example.com/e"])    # 非法字符
    ws.append([None, None])                          # 空白行，应跳过
    wb.save(path)

def make_second_sheet_roster(path):
    """验收：数据放在第二个工作表，仍只读取第一个。"""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["CodeID", "Content"])
    ws1.append(["Q001", "https://example.com/first"])
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["CodeID", "Content"])
    ws2.append(["Q100", "https://example.com/second"])
    wb.save(path)

def make_header_only_roster(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["CodeID", "Content"])
    wb.save(path)

def make_normal_csv(path):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        f.write("CodeID,Content\n")
        f.write("Q001,https://example.com/course/python\n")
        f.write("Q002,WIFI:T:WPA;S:TrainingRoom;P:12345678;;\n")
        f.write("Q003,mailto:training@example.com\n")
        f.write("Q004,TEL:+8613800000000\n")
        f.write("Q005,欢迎参加2026年秋季培训\n")

def main():
    base = Path(__file__).parent
    test_dir = base / "test_files"
    test_dir.mkdir(exist_ok=True)

    files = {
        "roster_normal.xlsx": make_normal_roster,
        "roster_abnormal.xlsx": make_abnormal_roster,
        "roster_second_sheet.xlsx": make_second_sheet_roster,
        "roster_header_only.xlsx": make_header_only_roster,
        "roster_normal.csv": make_normal_csv,
    }
    for name, fn in files.items():
        p = test_dir / name
        fn(p)
        print(f"已生成：{p}")

    # 5 张示例二维码 + manifest
    example_dir = base / "example_output"
    example_dir.mkdir(exist_ok=True)
    data = [
        ("Q001", "https://example.com/course/python"),
        ("Q002", "WIFI:T:WPA;S:TrainingRoom;P:12345678;;"),
        ("Q003", "mailto:training@example.com"),
        ("Q004", "TEL:+8613800000000"),
        ("Q005", "欢迎参加2026年秋季培训"),
    ]
    for cid, content in data:
        img = render_qr(content, 500)
        img.save(str(example_dir / f"{cid}.png"), "PNG")

    manifest_path = example_dir / "qr_manifest.csv"
    with open(manifest_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(
            f, fieldnames=["CodeID", "Content", "FileName", "Status"]
        )
        writer.writeheader()
        for cid, content in data:
            writer.writerow({
                "CodeID": cid,
                "Content": content,
                "FileName": f"{cid}.png",
                "Status": "Success",
            })
    print(f"已生成示例二维码目录：{example_dir}（5 张 PNG + qr_manifest.csv）")

if __name__ == "__main__":
    main()
